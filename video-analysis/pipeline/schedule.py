"""阶段 ①：读 Book1 时间表与 dta 人工标注，join 后校验，产出 segments。"""
from __future__ import annotations

import json
from dataclasses import dataclass, asdict
from pathlib import Path

import openpyxl
import pandas as pd

from pipeline import config


class ScheduleError(Exception):
    """输入数据不满足流水线假设时抛出。宁可中止，不可静默通过。"""


@dataclass(frozen=True)
class Presenter:
    seat_index: int
    gender_female: int
    speaks_pre: int
    speaks_qa: int


@dataclass(frozen=True)
class Segment:
    order: int
    company: str
    start: int
    end: int
    presenter_num: int
    presenters: tuple[Presenter, ...]
    dta_order: int = 0  # dta 中该企业的原始 order，供人工回溯核对；不用于下游编号
    # 会场最里侧是否坐着一位可见的专家（dta 的 Innermost_expert 列）。
    # 组级变量：同组内各企业取值相同。用户实测三组为 城市2=0、文化4=1、
    # 新能源2=0——与他逐企业核对时列出的一串「误把专家当创业者」完全对应
    # （那些全是文化4组）。默认 0 = 不额外剔除，是保守取值；默认 1 会在
    # 信息缺失时悄悄砍掉一个真人。
    innermost_expert: int = 0
    # 人工补录的实际进/退场时刻（dta 的 entryTime / exitTime），秒。
    # 有它就用它当进出场窗口的锚点（见 `windows.compute_windows`），没有才
    # 退回按段首/段尾猜。默认 None = 这一组还没人补，行为与补之前一字不差。
    entry_time: float | None = None
    exit_time: float | None = None


def _innermost_expert_of(row) -> int:
    """从 dta 行里取 Innermost_expert，缺失/NaN 一律记 0。

    记 0（不剔除）而非 1 是刻意的保守选择：这个变量缺失时如果默认「有专家」，
    流水线会在每段悄悄砍掉一个最里侧的人，而那很可能是后排的真实创业者——
    用户已明确指出「只拍到手不等于那里没人」，误删比漏判代价更大。
    """
    v = row.get("Innermost_expert") if hasattr(row, "get") else None
    try:
        if v is None or v != v:      # None 或 NaN
            return 0
        return 1 if int(float(v)) == 1 else 0
    except (TypeError, ValueError):
        return 0


def _mark_of(row, column: str) -> float | None:
    """从 dta 行里取 entryTime / exitTime，缺失（整列不存在 / NaN）返回 None。

    返回 None 而不是 0.0：0.0 是视频第 0 秒，是一个**合法**的时刻，用它当
    「没有」的哨兵，会让所有未补这两列的组把进出场窗口锚到片头去。
    """
    v = row.get(column) if hasattr(row, "get") else None
    try:
        if v is None or v != v:      # None 或 NaN
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def _read_xlsx(xlsx: Path, sheet: str) -> list[tuple[str, int, int]]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    if sheet not in wb.sheetnames:
        raise ScheduleError(
            f"Book1 中没有 sheet «{sheet}»。若这一组的 dta 里 startTime/"
            f"endTime 齐全，可用 `load_segments(..., from_dta=True)` "
            f"（config.SCHEDULE_FROM_DTA）改由 dta 单独排期，不要去改只读的输入文件")
    ws = wb[sheet]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        rows.append((str(r[0]).strip(), int(r[1]), int(r[2])))
    return rows


def _read_dta(dta: Path, panel: str) -> pd.DataFrame:
    df = pd.read_stata(dta, convert_categoricals=False)
    sub = df[df["panelName"].astype(str) == panel].copy()
    if sub.empty:
        raise ScheduleError(f"dta 中没有 panelName «{panel}»")
    return sub.sort_values("order").reset_index(drop=True)


def _int_or_zero(v) -> int:
    return 0 if pd.isna(v) else int(v)


def _segments_from_dta(dt: "pd.DataFrame") -> list[tuple[str, float, float]]:
    """dta 单独排期：直接用它的 order/firmName/startTime/endTime 当时间表。

    用户 2026-09-20 要处理电子10组——Book1 有 60 多组的 sheet，唯独没有这一组
    （电子9、电子11 都在），而 dta 里这 17 行的起止时刻一个不缺。为一组视频
    手工往只读的输入文件里补一张表，既动了「输入文件只读」这条底线，也没必要。

    段号取 dta 的 `order`，因此它必须是 1..N 连续：段号是
    `frames/{order:02d}`、`seg{order:02d}_seat{k}` 一长串路径的来源，跳号会让
    证据图与帧目录对不上，而表现只是「某一段的图找不到」。
    """
    orders = [int(v) for v in dt["order"].tolist()]
    if orders != list(range(1, len(orders) + 1)):
        raise ScheduleError(
            f"dta 单独排期要求 order 为 1..N 连续（段号直接取自它），实得 {orders}")
    rows: list[tuple[str, float, float]] = []
    for _, row in dt.iterrows():
        for col in ("startTime", "endTime"):
            if col not in row or pd.isna(row[col]):
                raise ScheduleError(
                    f"dta 单独排期需要 {col}，企业「{str(row['firmName']).strip()}」"
                    f"（order {int(row['order'])}）这一列是空的")
        rows.append((str(row["firmName"]).strip(),
                     float(row["startTime"]), float(row["endTime"])))
    return rows


def load_segments(
    xlsx: Path = config.XLSX,
    sheet: str = config.SHEET,
    dta: Path = config.DTA,
    panel: str = config.PANEL,
    from_dta: bool | None = None,
) -> list[Segment]:
    # `from_dta`：Book1 里没有这一组时，改用 dta 的 startTime/endTime 排期
    # （见 `_segments_from_dta`）。**显式开关，不做「读不到就回退」**——
    # 那样一个 sheet 名打错就会悄无声息地换一套时间表，而两套时间表给出的
    # 段号未必一致，错位的表现是整段的人全错、且一处不报错。
    if from_dta is None:
        from_dta = getattr(config, "SCHEDULE_FROM_DTA", False)
    dt = _read_dta(dta, panel)
    xl = _segments_from_dta(dt) if from_dta else _read_xlsx(xlsx, sheet)

    # join 语义：Book1 是 dta 的子集，按企业名匹配——不能按行号 zip，两边顺序
    # 不保证一致（详见 Task 27）。同名重复记录无法确定用哪条，硬失败。
    firm_names = dt["firmName"].map(lambda x: str(x).strip())
    dup_mask = firm_names.duplicated(keep=False)
    if dup_mask.any():
        dup_names = sorted(set(firm_names[dup_mask]))
        raise ScheduleError(f"dta 中企业名重复，无法确定对应记录：{'、'.join(dup_names)}")

    dta_by_name = {name: row for name, (_, row) in zip(firm_names, dt.iterrows())}

    segments: list[Segment] = []
    matched_names: set[str] = set()
    for i, (company, start, end) in enumerate(xl, start=1):
        row = dta_by_name.get(company)
        if row is None:
            raise ScheduleError(f"第 {i} 行企业「{company}」在 dta 中找不到人工标注")
        matched_names.add(company)
        if start >= end:
            raise ScheduleError(f"第 {i} 行起止时间非法：{start} >= {end}")

        n = int(row["presenterNum"])
        presenters = tuple(
            Presenter(
                seat_index=k,
                gender_female=_int_or_zero(row[f"prst{k}Female"]),
                speaks_pre=_int_or_zero(row[f"preprst{k}"]),
                speaks_qa=_int_or_zero(row[f"qaprst{k}"]),
            )
            for k in range(1, n + 1)
        )
        # 校验：非空 prstNFemale 的个数应与 presenterNum 一致
        non_null = sum(1 for k in range(1, 5) if not pd.isna(row[f"prst{k}Female"]))
        if non_null != n:
            raise ScheduleError(
                f"第 {i} 行 presenterNum={n}，但非空 prstNFemale 有 {non_null} 个"
            )

        segments.append(
            Segment(
                order=i,  # Book1 行序——下游 seg{order:02d}/frames/{order:02d} 等全部依赖它
                company=company,
                start=start,
                end=end,
                presenter_num=n,
                presenters=presenters,
                dta_order=int(row["order"]),
                innermost_expert=_innermost_expert_of(row),
                entry_time=_mark_of(row, "entryTime"),
                exit_time=_mark_of(row, "exitTime"),
            )
        )

    # dta 比 Book1 多出的企业不是错误（用户会之后补时间戳），但绝不能静默丢弃——
    # 必须显式打印，让人知道哪些企业当前跑不到。
    extra_names = [n for n in dta_by_name if n not in matched_names]
    if extra_names:
        extra_names.sort(key=lambda n: dta_by_name[n]["order"])
        parts = "、".join(
            f"{n}（标注 {int(dta_by_name[n]['presenterNum'])} 人）" for n in extra_names
        )
        print(f"⚠ dta 有 {len(extra_names)} 家企业在 Book1 中无时间戳，已排除：{parts}")

    for a, b in zip(segments, segments[1:]):
        if a.end > b.start:
            raise ScheduleError(f"order {a.order} 与 {b.order} 时间重叠")

    _check_marks(segments, panel=panel)

    total = sum(s.presenter_num for s in segments)
    if config.EXPECTED_PRESENTER_TOTAL == 0:
        print(f"ℹ EXPECTED_PRESENTER_TOTAL 为占位 0，跳过总人数断言；实际读到创业者总数 {total}，请回填 config.py")
    elif total != config.EXPECTED_PRESENTER_TOTAL:
        raise ScheduleError(
            f"创业者总数 {total} 与预期 {config.EXPECTED_PRESENTER_TOTAL} 不符"
        )

    return segments


def _check_marks(segments: list[Segment], panel: str | None = None,
                 confirmed=None) -> None:
    """校验人工补录的进/退场时刻落在合理区间内，否则硬失败。

    拦的是单位/数量级笔误（970 敲成 9700、分钟当成秒）。后果很具体：抽帧
    窗口锚在错的时刻上，抽出来的画面属于**别家企业**，却写进本段的
    `frames/NN/entry/` 目录，此后衣着/颜色/性别全部取自隔壁公司——而整条
    流水线照常跑完，一处不报错（`frames.segment_frame_mismatches` 的
    docstring 记的就是这类事故，那次是肉眼看出来的）。

    宽严的分寸：
      · 进场必须早于本段开始——晚于它就不是「进场」了。
      · 但**允许早到上一段还没讲完**：人在门口候场是真实会发生的，窗口本身
        会被 `compute_windows` 裁到上一段结束，不会真的抽到别人的画面。
        只有早过上一段**开始**才判定为笔误。
      · 退场：不得晚过下一段结束；早的那一侧留 `EXIT_MARK_EARLY_TOL` 的余地
        ——人起身时主持人还在收尾，全库 110 条里有 6 条只早了 0.4-5.5 秒，
        而真正的笔误早了 881 秒（见那个常数的注释）。

    **人工确认清单**（`config.MARKS_CONFIRMED`，用户 2026-09-27）：用户核对过
    视频、确认无误的「看似越界」的时刻，逐条放行并打印一行。键是
    `(组, 段号, "entry"|"exit", 数值)`——**连数值一起**：dta 里那一格日后被改
    成别的值，确认随之失效，一条旧确认不会替一个新笔误开绿灯。校验本身不关：
    它拦的是 970 敲成 9700 这类笔误，别的组照样需要。
    """
    if confirmed is None:
        confirmed = getattr(config, "MARKS_CONFIRMED", set())

    def _ok_by_human(seg, kind, value) -> bool:
        if panel is None or value is None:
            return False
        if (panel, int(seg.order), kind, round(float(value), 1)) in {
                (c[0], int(c[1]), c[2], round(float(c[3]), 1)) for c in confirmed}:
            print(f"   · order {seg.order}「{seg.company}」的"
                  f"{'进场' if kind == 'entry' else '退场'}时刻 {value} 超出常规范围，"
                  f"但已人工确认无误（config.MARKS_CONFIRMED），照常使用")
            return True
        return False

    starts = {s.order: s.start for s in segments}
    order_list = [s.order for s in segments]
    for i, seg in enumerate(segments):
        if seg.entry_time is not None:
            lo = starts[order_list[i - 1]] if i > 0 else 0.0
            if not (lo <= seg.entry_time <= seg.start) \
                    and not _ok_by_human(seg, "entry", seg.entry_time):
                raise ScheduleError(
                    f"order {seg.order}「{seg.company}」的进场时刻 "
                    f"{seg.entry_time} 不在 [{lo}, {seg.start}] 内——"
                    f"进场必须早于本段开始、且不得早过上一段开始，"
                    f"疑似 dta 的 entryTime 填错（按它抽帧会抽到别家企业的画面）")
        if seg.exit_time is not None:
            hi = segments[i + 1].end if i + 1 < len(segments) else float("inf")
            lo = seg.end - config.EXIT_MARK_EARLY_TOL
            if not (lo <= seg.exit_time <= hi) \
                    and not _ok_by_human(seg, "exit", seg.exit_time):
                raise ScheduleError(
                    f"order {seg.order}「{seg.company}」的退场时刻 "
                    f"{seg.exit_time} 不在 [{lo}, {hi}] 内——"
                    f"退场最多比本段结束早 {config.EXIT_MARK_EARLY_TOL:.0f} 秒、"
                    f"且不得晚过下一段结束，"
                    f"疑似 dta 的 exitTime 填错（按它抽帧会抽到别家企业的画面）")


def save_segments(segments: list[Segment], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "video": config.VIDEO.name,
        "sheet": config.SHEET,
        "panel": config.PANEL,
        "segments": [asdict(s) for s in segments],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_segments_json(path: Path) -> list[Segment]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    return [
        Segment(
            order=s["order"],
            company=s["company"],
            start=s["start"],
            end=s["end"],
            presenter_num=s["presenter_num"],
            presenters=tuple(Presenter(**p) for p in s["presenters"]),
            dta_order=s.get("dta_order", 0),
            innermost_expert=int(s.get("innermost_expert", 0) or 0),
            # 已跑完的三组的 segments.json 里没有这两个键 → None（退回按
            # 段首/段尾猜窗口），与它们此前的行为一字不差。
            entry_time=(None if s.get("entry_time") is None
                        else float(s["entry_time"])),
            exit_time=(None if s.get("exit_time") is None
                       else float(s["exit_time"])),
        )
        for s in payload["segments"]
    ]
